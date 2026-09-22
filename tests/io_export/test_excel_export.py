from pathlib import Path

import openpyxl

from plot_digitizer.calibration.axis import AxisCalibration, AxisScale
from plot_digitizer.calibration.transform import CoordinateTransform
from plot_digitizer.io_export.excel_export import export_curves_excel
from plot_digitizer.model.curve import Curve
from plot_digitizer.model.point import Point


def _identity_transform() -> CoordinateTransform:
    x_axis = AxisCalibration(scale=AxisScale.LINEAR)
    x_axis.add_reference(pixel=0.0, value=0.0)
    x_axis.add_reference(pixel=100.0, value=100.0)
    y_axis = AxisCalibration(scale=AxisScale.LINEAR)
    y_axis.add_reference(pixel=0.0, value=0.0)
    y_axis.add_reference(pixel=100.0, value=100.0)
    return CoordinateTransform(x_axis=x_axis, y_axis=y_axis)


def test_export_curves_excel_writes_one_sheet_per_curve(tmp_path: Path) -> None:
    transform = _identity_transform()
    curve_a = Curve(name="Alpha", points=[Point(x=0.0, y=0.0), Point(x=10.0, y=10.0)])
    curve_b = Curve(name="Beta", points=[Point(x=5.0, y=5.0)])

    out_path = tmp_path / "all_curves.xlsx"
    export_curves_excel([curve_a, curve_b], transform, out_path)

    workbook = openpyxl.load_workbook(out_path)
    assert workbook.sheetnames == ["Alpha", "Beta"]
    alpha_rows = list(workbook["Alpha"].iter_rows(values_only=True))
    assert alpha_rows == [("x", "y"), (0.0, 0.0), (10.0, 10.0)]
    beta_rows = list(workbook["Beta"].iter_rows(values_only=True))
    assert beta_rows == [("x", "y"), (5.0, 5.0)]


def test_export_curves_excel_sanitizes_invalid_sheet_name_characters(tmp_path: Path) -> None:
    transform = _identity_transform()
    curve = Curve(name="Bad: [Name]/Test?*", points=[Point(x=1.0, y=1.0)])

    out_path = tmp_path / "sanitized.xlsx"
    export_curves_excel([curve], transform, out_path)

    workbook = openpyxl.load_workbook(out_path)
    assert workbook.sheetnames == ["Bad NameTest"]


def test_export_curves_excel_deduplicates_identical_sheet_names(tmp_path: Path) -> None:
    transform = _identity_transform()
    curve_1 = Curve(name="Curve", points=[Point(x=0.0, y=0.0)])
    curve_2 = Curve(name="Curve", points=[Point(x=1.0, y=1.0)])

    out_path = tmp_path / "duplicates.xlsx"
    export_curves_excel([curve_1, curve_2], transform, out_path)

    workbook = openpyxl.load_workbook(out_path)
    assert workbook.sheetnames == ["Curve", "Curve_2"]


def test_export_curves_excel_truncates_long_sheet_names(tmp_path: Path) -> None:
    transform = _identity_transform()
    long_name = "A" * 50
    curve = Curve(name=long_name, points=[Point(x=0.0, y=0.0)])

    out_path = tmp_path / "long_name.xlsx"
    export_curves_excel([curve], transform, out_path)

    workbook = openpyxl.load_workbook(out_path)
    assert len(workbook.sheetnames[0]) <= 31
